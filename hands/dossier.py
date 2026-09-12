"""Complaint dossier builder — 8 sections, R10 completeness-gated.

The state-ladder differentiator: company bot → supervisor →
@consumergovuz_bot / 1159. The dossier travels with the user and is
exported with real corpus citations (source URL + fetch date in each).
"""
from __future__ import annotations

import os

from docx import Document

SECTIONS = (
    "1. Citizen (Aholi)",
    "2. Counterparty (Karshi tomon)",
    "3. Order reference (Buyurtma)",
    "4. Narrative (Voqealar bayoni)",
    "5. Legal basis (Huquqiy asos)",
    "6. Counterparty policy (Karshi tomon siyosati)",
    "7. Evidence (Islotlar ro'yxati)",
    "8. Demand & deadline (Talab va muddat)",
)


def completeness(report: dict) -> tuple[int, list[str]]:
    """(n_present, missing) — R10: export only when nothing is missing."""
    missing = [s for s in SECTIONS if not str(report.get(s, "")).strip()]
    return len(SECTIONS) - len(missing), missing


def build(case: dict, report: dict, outdir: str | None = None) -> str:
    """case: {id, mandate}; report: keys = SECTIONS. Returns the .docx path."""
    _, missing = completeness(report)
    if missing:
        raise ValueError(f"R10-dossier-completeness: missing sections: {missing}")

    doc = Document()
    doc.add_heading("ARIZA / COMPLAINT DOSSIER", 0)
    doc.add_paragraph(f"Case: {case.get('id')}   ·   Objective: {case.get('mandate', {}).get('objective', '')}")
    for s in SECTIONS:
        doc.add_heading(s, level=1)
        doc.add_paragraph(str(report[s]))
    doc.add_paragraph("— wakil avtomatik tayyorladi; hamma iqtibos manbali (korpus, sana bilan).")

    outdir = outdir or os.path.join(os.path.dirname(__file__), "outputs")
    os.makedirs(outdir, exist_ok=True)
    path = os.path.join(outdir, f"dossier-{case.get('id', 'case')}.docx")
    doc.save(path)
    return path
